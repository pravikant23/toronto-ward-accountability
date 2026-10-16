#!/usr/bin/env python3
"""
data/validate.py
Hard-fail validation of data/wards.json against source files.
The build must FAIL if any check fails.
Run after build_data.py.
"""

import json
import os
import sys
import random
import openpyxl
from collections import Counter

BASE     = os.path.dirname(os.path.abspath(__file__))
RAW      = os.path.join(BASE, "raw")
OUT_JSON = os.path.join(BASE, "wards.json")

VOTER_2022   = os.path.join(RAW, "2022_voter_turnout_statistics_final.xlsx")
RESULTS_2022 = os.path.join(RAW, "2022-results", "2022_Toronto_Poll_By_Poll_Councillor.xlsx")

PASS_COUNT = 0
FAIL_COUNT = 0


def ok(msg):
    global PASS_COUNT
    PASS_COUNT += 1
    print(f"  PASS  {msg}")


def fail(msg):
    global FAIL_COUNT
    FAIL_COUNT += 1
    print(f"  FAIL  {msg}", file=sys.stderr)


# ── Load wards.json ────────────────────────────────────────────────────────────
with open(OUT_JSON, encoding="utf-8") as f:
    data = json.load(f)

wards = data["wards"]

print("=== Structural checks ===")

# 1. Exactly 25 wards
if len(wards) == 25:
    ok(f"Exactly 25 wards found")
else:
    fail(f"Expected 25 wards, got {len(wards)}")

# 2. Ward numbers 1–25 with no duplicates
ward_nums = [w["ward_num"] for w in wards]
if sorted(ward_nums) == list(range(1, 26)):
    ok("Ward numbers 1–25, no duplicates")
else:
    missing  = sorted(set(range(1, 26)) - set(ward_nums))
    extra    = sorted(set(ward_nums) - set(range(1, 26)))
    fail(f"Ward numbers wrong — missing: {missing}, extra: {extra}")

# 3. Exactly one winner per ward
print("\n=== Winner checks ===")
for w in wards:
    winners = [c for c in w["candidates_2022"] if c.get("is_winner")]
    if len(winners) == 1:
        ok(f"Ward {w['ward_num']:>2}: exactly one winner ({winners[0]['name']})")
    else:
        fail(f"Ward {w['ward_num']:>2}: {len(winners)} winners")

# 4. Candidate vote sum ≤ ballots_cast
print("\n=== Vote count checks ===")
for w in wards:
    total_cand_votes = sum(c["votes"] for c in w["candidates_2022"])
    ballots          = w["ballots_cast"] or 0
    if total_cand_votes <= ballots:
        ok(
            f"Ward {w['ward_num']:>2}: candidate votes ({total_cand_votes:,}) "
            f"≤ ballots cast ({ballots:,})"
        )
    else:
        fail(
            f"Ward {w['ward_num']:>2}: candidate votes ({total_cand_votes:,}) "
            f"> ballots cast ({ballots:,})"
        )

# 5. Our computed turnout matches city's published turnout within 0.1 pt
print("\n=== Turnout accuracy checks ===")
for w in wards:
    reg = w["registered_voters"]
    bc  = w["ballots_cast"]
    tp  = w["turnout_pct"]
    if reg and bc and tp is not None:
        computed = round(bc / reg * 100, 2)
        diff     = abs(computed - tp)
        if diff <= 0.1:
            ok(f"Ward {w['ward_num']:>2}: turnout {tp}% matches computed {computed}% (diff={diff:.2f})")
        else:
            fail(
                f"Ward {w['ward_num']:>2}: turnout {tp}% ≠ computed {computed}% "
                f"(diff={diff:.4f})"
            )
    else:
        fail(f"Ward {w['ward_num']:>2}: missing registration or ballots data")

# 6. City-level turnout = 29.66% (Number Voted only, not including declined)
print("\n=== City-level check ===")
city = data.get("city_2022", {})
city_tp = city.get("turnout_pct")
if city_tp == 29.66:
    ok(f"City 2022 turnout = {city_tp}% (matches expected 29.66%)")
else:
    fail(f"City 2022 turnout = {city_tp}%, expected 29.66%")

city_reg = city.get("registered_voters")
city_bc  = city.get("ballots_cast")
if city_reg == 1_898_750 and city_bc == 563_124:
    ok(f"City 2022 registered={city_reg:,}, ballots={city_bc:,} match expected values")
else:
    fail(
        f"City 2022 totals: registered={city_reg:,} (exp 1,898,750), "
        f"ballots={city_bc:,} (exp 563,124)"
    )

# 7. No null/zero placeholders for critical fields
print("\n=== Null/zero checks ===")
REQUIRED_NONZERO = [
    "registered_voters",
    "ballots_cast",
    "turnout_pct",
    "turnout_2018_pct",
    "winner_votes",
    "num_candidates",
]
for w in wards:
    for field in REQUIRED_NONZERO:
        val = w.get(field)
        if val is None:
            fail(f"Ward {w['ward_num']:>2}: {field} is None")
        elif val == 0:
            fail(f"Ward {w['ward_num']:>2}: {field} is zero")
        else:
            ok(f"Ward {w['ward_num']:>2}: {field} = {val}")

# 8. Census data not null for each ward
print("\n=== Census null checks ===")
CENSUS_REQUIRED = [
    "population_2021",
    "median_age_2021",
    "pct_owned",
    "pct_rented",
    "median_household_income_2020",
]
for w in wards:
    cen = w.get("census", {})
    for field in CENSUS_REQUIRED:
        val = cen.get(field)
        if val is None:
            fail(f"Ward {w['ward_num']:>2}: census.{field} is None")
        else:
            ok(f"Ward {w['ward_num']:>2}: census.{field} = {val}")

# 9. By-election wards have by_election note; others don't
print("\n=== By-election flag checks ===")
BY_ELECTION_WARDS = {15, 20, 25}
for w in wards:
    wn  = w["ward_num"]
    bye = w.get("by_election")
    if wn in BY_ELECTION_WARDS:
        if bye:
            ok(f"Ward {wn:>2}: has by_election note")
        else:
            fail(f"Ward {wn:>2}: missing by_election note (expected one)")
    else:
        if bye is None:
            ok(f"Ward {wn:>2}: no by_election note (correct)")
        else:
            fail(f"Ward {wn:>2}: unexpected by_election note: {bye}")

# 10. Spot-check 3 random wards against raw source files
print("\n=== Spot-check (3 random wards vs raw source) ===")
random.seed(42)
spot_wards = random.sample(range(1, 26), 3)

# Load raw voter 2022 for spot check
wb_vt = openpyxl.load_workbook(VOTER_2022, read_only=True, data_only=True)
ws_vt = wb_vt["2022 Voter Turnout Statistics"]
raw_eligible = {}
raw_voted    = {}
raw_declined = {}
current_ward = None
for i, row in enumerate(ws_vt.iter_rows(values_only=True)):
    if i == 0:
        continue
    wc = row[0]
    sc = row[1]
    if isinstance(wc, int):
        current_ward = wc
    if current_ward is None:
        continue
    try:
        sub_num = int(sc)
    except (TypeError, ValueError):
        continue
    if 1 <= sub_num <= 96:
        raw_eligible[current_ward] = raw_eligible.get(current_ward, 0) + (row[13] or 0)
    raw_voted[current_ward]    = raw_voted.get(current_ward, 0)    + (row[14] or 0)
    raw_declined[current_ward] = raw_declined.get(current_ward, 0) + (row[17] or 0)

# Load raw results 2022 for spot check
wb_res = openpyxl.load_workbook(RESULTS_2022, read_only=True, data_only=True)
raw_winner_votes = {}
for sheet_name in wb_res.sheetnames:
    if not sheet_name.startswith("Ward "):
        continue
    try:
        wn = int(sheet_name.split()[1])
    except (IndexError, ValueError):
        continue
    ws_r = wb_res[sheet_name]
    rows = list(ws_r.iter_rows(values_only=True))
    cands = []
    for row in rows[3:]:
        nc, vc = row[0], row[-1]
        if nc and vc and isinstance(vc, (int, float)):
            ns = str(nc).strip()
            if ns and "Totals" not in ns:
                cands.append(int(vc))
    if cands:
        raw_winner_votes[wn] = max(cands)

for wn in spot_wards:
    wd      = next(w for w in wards if w["ward_num"] == wn)
    exp_reg = int(raw_eligible.get(wn, 0))
    exp_bc  = int(raw_voted.get(wn, 0))  # Number Voted only (not + declined)
    exp_wv  = raw_winner_votes.get(wn)

    if wd["registered_voters"] == exp_reg:
        ok(f"Ward {wn:>2}: registered_voters {exp_reg:,} matches raw file")
    else:
        fail(
            f"Ward {wn:>2}: registered_voters {wd['registered_voters']:,} "
            f"≠ raw {exp_reg:,}"
        )

    if wd["ballots_cast"] == exp_bc:
        ok(f"Ward {wn:>2}: ballots_cast {exp_bc:,} matches raw file (Number Voted only)")
    else:
        fail(
            f"Ward {wn:>2}: ballots_cast {wd['ballots_cast']:,} "
            f"≠ raw {exp_bc:,} (Number Voted only)"
        )

    if exp_wv is not None and wd["winner_votes"] == exp_wv:
        ok(f"Ward {wn:>2}: winner_votes {exp_wv:,} matches raw file")
    elif exp_wv is not None:
        fail(
            f"Ward {wn:>2}: winner_votes {wd['winner_votes']:,} "
            f"≠ raw {exp_wv:,}"
        )

# 11. Accuracy fixes — data-source and period-label checks
print("\n=== Accuracy checks (period labels, language source, ballots) ===")

# 11a. All 25 wards must have exactly 3 top languages; each must have a positive speaker count
for w in wards:
    langs = (w.get("census") or {}).get("top_non_official_languages", [])
    if len(langs) == 3:
        ok(f"Ward {w['ward_num']:>2}: 3 top languages present")
    else:
        fail(f"Ward {w['ward_num']:>2}: expected 3 top languages, got {len(langs)}")
    for lg in langs:
        name = lg.get("language", "")
        spk  = lg.get("speakers", 0)
        if name.lower() in ("english", "french"):
            fail(f"Ward {w['ward_num']:>2}: language list contains '{name}' (must be excluded)")
        if "total" in name.lower() or "non-official" in name.lower():
            fail(f"Ward {w['ward_num']:>2}: language '{name}' looks like an aggregate row (must be excluded)")
        if not isinstance(spk, int) or spk <= 0:
            fail(f"Ward {w['ward_num']:>2}: language '{name}' has invalid speaker count {spk}")
        else:
            ok(f"Ward {w['ward_num']:>2}: language '{name}' = {spk:,} speakers")

# 11b. Ward 4 language spot-check (Tibetan, Polish, Spanish)
w4 = next((w for w in wards if w["ward_num"] == 4), None)
if w4:
    langs4 = {lg["language"]: lg["speakers"] for lg in
              (w4.get("census") or {}).get("top_non_official_languages", [])}
    for name, expected in [("Tibetan", 1330), ("Polish", 1210), ("Spanish", 1205)]:
        got = langs4.get(name)
        if got == expected:
            ok(f"Ward  4: {name} = {got:,} (matches expected)")
        else:
            fail(f"Ward  4: {name} = {got} (expected {expected})")

# 11b-2. Ward 4 turnout spot-check: 2022 and 2018 values traced to raw xlsx files
w4_tp22 = w4.get("turnout_pct") if w4 else None
w4_tp18 = w4.get("turnout_2018_pct") if w4 else None
if w4_tp22 == 38.45:
    ok("Ward  4: turnout_2022_pct = 38.45% (confirmed against raw xlsx)")
else:
    fail(f"Ward  4: turnout_2022_pct = {w4_tp22}, expected 38.45%")
if w4_tp18 == 46.53:
    ok("Ward  4: turnout_2018_pct = 46.53% (confirmed against raw xlsx, NOT 46.57)")
else:
    fail(f"Ward  4: turnout_2018_pct = {w4_tp18}, expected 46.53%")

# 11c. City-wide ballots = 563,124 (Number Voted only, not including declined)
bc_city = city_bc
if bc_city == 563_124:
    ok(f"City ballots_cast = {bc_city:,} (Number Voted only, correct)")
else:
    fail(f"City ballots_cast = {bc_city:,}, expected 563,124 (Number Voted only)")

# 11d. City-wide turnout = 29.66%
if city_tp == 29.66:
    ok(f"City turnout = {city_tp}% (Number Voted ÷ registered, correct)")
else:
    fail(f"City turnout = {city_tp}%, expected 29.66%")

# 11e. Issues data present: all 25 wards must have crime.ytd_jan_sep_2024
for w in wards:
    ytd24 = ((w.get("issues") or {}).get("crime") or {}).get("ytd_jan_sep_2024")
    if ytd24 and ytd24.get("count") is not None:
        ok(f"Ward {w['ward_num']:>2}: crime.ytd_jan_sep_2024.count = {ytd24['count']:,}")
    else:
        fail(f"Ward {w['ward_num']:>2}: crime.ytd_jan_sep_2024 missing or null")

# 11f. City Jan-Sep crime 2024 = 35,378 and 2025 = 30,688 (spot check via Ward 1)
w1 = next((w for w in wards if w["ward_num"] == 1), None)
if w1:
    c1 = (w1.get("issues") or {}).get("crime") or {}
    city_24 = (c1.get("city_ytd_jan_sep_2024") or {}).get("count")
    city_25 = (c1.get("city_jan_sep_2025") or {}).get("count")
    if city_24 == 35_378:
        ok(f"City Jan-Sep 2024 crime = {city_24:,} (matches expected 35,378)")
    else:
        fail(f"City Jan-Sep 2024 crime = {city_24}, expected 35,378")
    if city_25 == 30_688:
        ok(f"City Jan-Sep 2025 crime = {city_25:,} (matches expected 30,688)")
    else:
        fail(f"City Jan-Sep 2025 crime = {city_25}, expected 30,688")

# 11g. Budget votes must have vote_description stored for all years
for w in wards:
    for bv in (w.get("budget_votes") or []):
        vd = bv.get("vote_description")
        if vd and len(vd) > 2:
            ok(f"Ward {w['ward_num']:>2}: budget {bv['year']} vote_description present: {vd[:40]}")
        else:
            fail(f"Ward {w['ward_num']:>2}: budget {bv['year']} vote_description missing or blank")

# 11h. Close votes must not include known procedural motion types
FORBIDDEN_PROCEDURAL = {"Waive Referral", "End Debate", "Re-open Item",
                        "Waive Notice", "Introduce Motion without Notice"}
cv_all_types: Counter = Counter()
for w in wards:
    for cv in (w.get("close_votes") or []):
        cv_all_types[cv.get("motion_type", "")] += 1
for mt in FORBIDDEN_PROCEDURAL:
    if cv_all_types[mt] > 0:
        fail(f"Close votes contain procedural motion type '{mt}' ({cv_all_types[mt]} occurrences)")
    else:
        ok(f"Close votes: '{mt}' correctly excluded")

# 12. Councillor metrics and detail-page checks
import csv as _csv

METRICS_CSV  = os.path.join(BASE, "councillor_metrics.csv")
SUMMARIES_CSV = os.path.join(BASE, "vote_summaries.csv")

print("\n=== Councillor metrics checks ===")

# 12a. councillor_metrics.csv — test answers and n-for-every-pct
if not os.path.exists(METRICS_CSV):
    fail("councillor_metrics.csv not found — run data/councillor_metrics.py first")
else:
    ok("councillor_metrics.csv found")
    _metrics = {}
    with open(METRICS_CSV, newline="", encoding="utf-8") as _f:
        for _row in _csv.DictReader(_f):
            for _k in ("with_majority_pct", "with_mayor_pct", "attendance_pct",
                       "contested_cast", "majority_n", "mayor_n",
                       "sessions_present", "sessions_total"):
                _v = _row.get(_k, "")
                try:
                    _row[_k] = float(_v) if "." in (_v or "") else (int(_v) if _v else None)
                except ValueError:
                    _row[_k] = None
            _metrics[_row["fullname"]] = _row

    _TEST_MEMBERS = [
        ("Gord Perks",      185, 80.5, 85.2),
        ("Stephen Holyday", 191, 30.4, 22.7),
        ("Amber Morley",    187, 93.0, 95.2),
        ("Neethan Shan",     49, 91.8, 89.1),
    ]
    for _name, _exp_cast, _exp_maj, _exp_may in _TEST_MEMBERS:
        _m = _metrics.get(_name)
        if _m is None:
            fail(f"{_name}: not in councillor_metrics.csv")
        else:
            _ok = (_m["contested_cast"] == _exp_cast and
                   _m["with_majority_pct"] == _exp_maj and
                   _m["with_mayor_pct"] == _exp_may)
            if _ok:
                ok(f"{_name}: contested_cast={_exp_cast}, "
                   f"with_majority={_exp_maj}%, with_mayor={_exp_may}%")
            else:
                fail(f"{_name}: expected cast={_exp_cast}/{_exp_maj}%/{_exp_may}%, "
                     f"got {_m['contested_cast']}/{_m['with_majority_pct']}%/{_m['with_mayor_pct']}%")

    # Every pct has its n
    for _fn, _m in _metrics.items():
        if _m.get("with_majority_pct") is not None:
            _n = int(_m.get("majority_n") or 0)
            if _n > 0:
                ok(f"{_fn}: with_majority_pct has n={_n}")
            else:
                fail(f"{_fn}: with_majority_pct set but majority_n=0")
        if _m.get("with_mayor_pct") is not None:
            _n = int(_m.get("mayor_n") or 0)
            if _n > 0:
                ok(f"{_fn}: with_mayor_pct has n={_n}")
            else:
                fail(f"{_fn}: with_mayor_pct set but mayor_n=0")
        if _m.get("attendance_pct") is not None:
            _pre = int(_m.get("sessions_present") or 0)
            _tot = int(_m.get("sessions_total") or 0)
            if _tot > 0:
                ok(f"{_fn}: attendance_pct has n={_tot}")
            else:
                fail(f"{_fn}: attendance_pct set but sessions_total=0")

print("\n=== Vote summaries checks ===")

# 12b. vote_summaries.csv — banned words, official URLs, all items covered
_BANNED = {"attempt", "blocked", "pushed", "controversial", "rejected", "failed"}

if not os.path.exists(SUMMARIES_CSV):
    fail("vote_summaries.csv not found — run data/generate_summaries.py first")
else:
    ok("vote_summaries.csv found")
    _summaries = {}
    with open(SUMMARIES_CSV, newline="", encoding="utf-8") as _f:
        for _row in _csv.DictReader(_f):
            _summaries[_row["item_num"]] = _row

    for _item, _s in _summaries.items():
        _summary_lc = (_s.get("summary") or "").lower()
        _bad = [_w for _w in _BANNED if _w in _summary_lc]
        if _bad:
            fail(f"{_item}: summary contains banned word(s): {_bad}")
        else:
            ok(f"{_item}: no banned words in summary")

        _url = _s.get("official_url", "")
        if _url.startswith("https://secure.toronto.ca/council/agenda-item.do?item="):
            ok(f"{_item}: has official TMMIS URL")
        else:
            fail(f"{_item}: missing or invalid official_url: {_url!r}")

    # Every close-vote item must have a summary entry
    _missing = set()
    for w in wards:
        for cv in (w.get("close_votes") or []):
            _it = cv.get("item", "")
            if _it and _it not in _summaries:
                _missing.add(_it)
    if _missing:
        for _it in sorted(_missing):
            fail(f"Close-vote item {_it} has no entry in vote_summaries.csv")
    else:
        ok("All close-vote items have a summary entry in vote_summaries.csv")

    # No close-vote title is truncated (must not end with '…')
    _truncated = []
    for w in wards:
        for cv in (w.get("close_votes") or []):
            if cv.get("title", "").endswith("…"):
                _truncated.append(cv["item"])
    if _truncated:
        for _it in _truncated:
            fail(f"Close-vote title for {_it} is truncated")
    else:
        ok("No close-vote titles are truncated")

# ── Final result ───────────────────────────────────────────────────────────────
print()
print(f"{'=' * 50}")
print(f"Results: {PASS_COUNT} passed, {FAIL_COUNT} failed")
print(f"{'=' * 50}")

if FAIL_COUNT > 0:
    print("\nBUILD FAILED — fix errors above before proceeding.", file=sys.stderr)
    sys.exit(1)
else:
    print("\nAll checks passed. wards.json is valid.")
